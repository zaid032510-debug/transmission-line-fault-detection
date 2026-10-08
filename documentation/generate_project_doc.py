from docx import Document


doc = Document()
doc.add_heading("Machine Learning-Based Transmission Line Fault Detection Using the IEEE 68-Bus System", level=1)

paragraphs = [
    "This project focuses on developing an automated transmission-line fault detection system using machine learning. The IEEE 68-bus test system is used as the power-system model, with bus and transmission-line information obtained from the available PSS/E .RAW and .DYR files.",
    "Since PSS/E is not installed, the project uses the open-source ANDES power-system simulation tool to load and simulate the IEEE 68-bus system. The simulation data will be used to generate normal operating conditions and transmission-line fault conditions.",
    "The project follows a complete pipeline:",
    "IEEE 68-Bus System → Power-System Simulation → Data Extraction → Feature Engineering → Fault Labeling → Machine Learning → Fault Detection",
    "The extracted electrical parameters, such as bus voltage magnitude, voltage angle, and other relevant system measurements, are processed into a machine-learning dataset. Fault cases will then be labeled according to their operating condition.",
    "Two machine-learning algorithms are planned for fault classification:",
    "Decision Tree",
    "Random Forest",
    "The trained models will identify whether the system is operating under a normal condition or a fault condition and can be extended to determine the faulted transmission-line condition.",
    "A dashboard is also included in the project structure to provide a user-friendly way of presenting the detection results.",
    "Main Objective",
    "The main objective is to develop a data-driven transmission-line fault detection system that can analyze electrical-system measurements and automatically detect abnormal/fault conditions, reducing the dependence on manual analysis and providing a foundation for faster power-system protection and monitoring.",
    "Current Project Status",
    "The initial data-processing pipeline is already working:",
]

for item in paragraphs:
    if item in {"Main Objective", "Current Project Status"}:
        doc.add_heading(item, level=2)
    else:
        doc.add_paragraph(item)

bullet_items = [
    "IEEE 68-bus RAW data → 68 bus records",
    "Feature generation → 68 feature records",
    "ML dataset generation → 68 records",
    "ANDES 2.0.0 has been installed successfully.",
    "The next step is to simulate the IEEE 68-bus system and generate actual normal and fault-condition data.",
    "One important point: the current fault = 0 assignment is only a temporary placeholder. It will be replaced with genuine fault labels after the simulation data is generated.",
    "Available next action: Create a downloadable DOCX file here in this chat containing the editable prose above. check its files",
]

for item in bullet_items:
    doc.add_paragraph(item, style="List Bullet")

output_path = r"e:\Transmission_Line_Fault_project\documentation\Final_Year_Project_Description.docx"
doc.save(output_path)
print(output_path)
